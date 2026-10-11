"""Звіт для власника бізнесу: що в ньому має бути й чого бути не може."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.audit.enums import Module, Severity
from app.audit.i18n import Locale
from app.audit.scoring import compute_scores
from app.audit.types import Issue
from app.report import build_report
from app.report.docx_owner import filename_for, render_owner_docx
from app.report.types import ReportSource


def _source(**overrides) -> ReportSource:
    issues = [
        Issue(
            code="seo.h1",
            module=Module.SEO,
            severity=Severity.CRITICAL,
            title="Немає H1",
            msg="seo.h1.missing",
            weight=4,
            page_url="https://example.com/",
        ),
        Issue(
            code="tech.analytics",
            module=Module.TECH,
            severity=Severity.HIGH,
            title="Немає аналітики",
            msg="tech.analytics.missing",
            weight=3,
        ),
    ]
    # Скор рахується з тих самих issue, що й потенціал: у справжньому
    # аудиті вони походять з одного джерела, і фікстура з розбіжними
    # числами перевіряла б неіснуючий стан.
    extra = [
        Issue(
            code="perf.core_web_vitals",
            module=Module.PERFORMANCE,
            severity=Severity.CRITICAL,
            title="Повільно",
            msg="perf.core_web_vitals.lcp",
            params={"value": 5000, "unit": "ms", "strategy": "mobile", "threshold": 2500},
            weight=4,
        ),
        Issue(
            code="ai.jsonld.missing",
            module=Module.AI_READY,
            severity=Severity.HIGH,
            title="Немає розмітки",
            msg="ai.jsonld.missing",
            weight=4,
            page_url="https://example.com/",
        ),
    ]
    issues = issues + extra
    overall, scores = compute_scores(issues)
    defaults = dict(
        url="https://example.com",
        entry_url="https://example.com/",
        origin="https://example.com",
        overall=overall,
        scores=scores,
        issues=issues,
        degraded=[],
        skipped_modules=set(),
        evidence={},
    )
    defaults.update(overrides)
    return ReportSource(**defaults)  # type: ignore[arg-type]


def _text(blob: bytes) -> str:
    document = Document(io.BytesIO(blob))
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_report_opens_as_a_valid_word_file():
    blob = render_owner_docx(build_report(_source(), Locale.UK))
    assert Document(io.BytesIO(blob)).paragraphs, "порожній документ клієнту не віддаси"


def test_owner_report_never_tells_how_to_fix_anything():
    """Порада «що зробити» — це те, що продається. У звіті вона віддає
    товар безкоштовно, а власник її все одно не виконуватиме сам."""
    report = build_report(_source(), Locale.UK)
    text = _text(render_owner_docx(report, Locale.UK))

    for action in (r.action for r in report.recommendations if r.action):
        assert action not in text, f"інструкція з технічного звіту протекла: {action}"


def test_owner_report_speaks_about_consequences():
    report = build_report(_source(), Locale.UK)
    text = _text(render_owner_docx(report, Locale.UK))

    consequences = [r.consequence for r in report.recommendations if r.consequence]
    assert any(c in text for c in consequences), "без наслідків документ ні про що"


def test_areas_are_named_the_way_an_owner_would_name_them():
    text = _text(render_owner_docx(build_report(_source(), Locale.UK)))

    assert "Чи знаходять вас у пошуку" in text
    for internal in ("seo", "ai_ready", "llm_visibility", "performance"):
        assert internal not in text, f"назва модуля протекла у звіт: {internal}"


def test_missing_overall_does_not_break_the_document():
    """Сайт віддав ботам 403: бала немає, документ має лишитися цілим."""
    blob = render_owner_docx(build_report(_source(overall=None), Locale.UK))

    assert "—" in _text(blob)


def test_potential_score_is_not_the_sum_of_separate_gains():
    """Сума impact_points завищена: стеля на code і нелінійність шкали.
    Обіцяти приріст, якого не буде, — найгірше для комерційного звіту."""
    report = build_report(_source(), Locale.UK)
    naive_sum = sum(a.expected_gain_points for a in report.next_actions)

    assert report.potential_score is not None
    assert report.overall_score is not None
    assert report.potential_score - report.overall_score <= naive_sum


def test_filename_names_the_kind_of_document():
    """Вид, а не аудиторія: власник затвердив «загальний» і «технічний»."""
    name = filename_for(build_report(_source(), Locale.UK), Locale.UK)

    assert "general" in name and name.endswith(".docx")
    assert "owner" not in name


def test_english_version_has_no_ukrainian_left():
    text = _text(render_owner_docx(build_report(_source(), Locale.EN), Locale.EN))

    assert "Чи знаходять" not in text
    assert "What we found" in text


# --- два режими: базовий і розширений -----------------------------------

def _report(evidence: dict | None = None):
    return build_report(_source(evidence=evidence or {}), Locale.UK)


def test_base_mode_never_calls_a_model(monkeypatch):
    """Базовий звіт має працювати завжди — зокрема без ключа й без мережі.
    Будь-який виклик тут означав би, що документ мовчки став платним."""
    def explode(*_args, **_kwargs):
        raise AssertionError("базовий режим пішов у мережу")

    monkeypatch.setattr("app.audit.llm.factory.get_analyst", explode, raising=False)

    blob = render_owner_docx(_report(), Locale.UK)

    assert len(blob) > 5000


def test_base_mode_says_nothing_about_enrichment():
    """Розширення не просили — про його відсутність не йдеться."""
    text = _text(render_owner_docx(_report(), Locale.UK))

    assert "Розгорнутих пояснень" not in text
    assert "штучний інтелект описує" not in text


def test_enriched_mode_adds_the_personal_wording():
    from app.report.owner_enrich import OwnerExtras

    report = _report()
    first = report.recommendations[0]
    extras = OwnerExtras(
        requested=True,
        analysed=True,
        opening="Ваша пекарня продає печиво, але сайт про це майже мовчить.",
        ai_view="Компанія випікає печиво та хліб.",
        notes={first.msg: "Покупець не зрозуміє з сайту, що ви печете саме на замовлення."},
    )

    text = _text(render_owner_docx(report, Locale.UK, extras))

    assert "Ваша пекарня продає печиво" in text
    assert "Компанія випікає печиво та хліб." in text
    assert "Покупець не зрозуміє з сайту" in text
    # Доведений текст із каталогу лишається: переказ додається, не замінює.
    assert first.consequence in text


def test_failed_enrichment_says_so_instead_of_looking_like_the_base_report():
    """Мовчазний відкат виглядав би так, ніби розширений звіт саме такий
    і має бути, — а користувач за нього заплатив."""
    from app.report.owner_enrich import OwnerExtras

    extras = OwnerExtras(requested=True, analysed=False, error="не заданий ключ")

    text = _text(render_owner_docx(_report(), Locale.UK, extras))

    assert "не заданий ключ" in text


def test_filenames_differ_so_the_two_documents_do_not_overwrite_each_other():
    from app.report.owner_enrich import OwnerExtras

    report = _report()
    base = filename_for(report, Locale.UK)
    rich = filename_for(report, Locale.UK, OwnerExtras(requested=True, analysed=True))

    assert base != rich


def test_ai_view_comes_from_the_stored_brief_without_a_new_call():
    """За brief уже заплачено під час аудиту. Другий виклик заради того
    самого тексту — це рахунок клієнту за те, що вже лежить у базі."""
    from app.report.owner_enrich import ai_view_of, brief_from_evidence

    report = _report({"site_brief": {"analysed": True, "description": "Пекарня в Києві."}})

    assert ai_view_of(brief_from_evidence(report.evidence)) == "Пекарня в Києві."


def test_unanalysed_brief_gives_no_ai_view():
    from app.report.owner_enrich import ai_view_of

    assert ai_view_of({"analysed": False, "error": "ключа немає"}) == ""
    assert ai_view_of({}) == ""


@pytest.mark.asyncio
async def test_extras_never_invent_a_problem_we_did_not_find():
    """Модель повертає ключ, якого ми не давали. У звіт він потрапити не
    може: це була б знахідка без доказу."""
    from app.audit.llm.base import OwnerNotes
    from app.report.owner_enrich import build_owner_extras

    class Inventive:
        name = "fake"

        async def owner_notes(self, business, band, problems):
            return OwnerNotes(
                analysed=True,
                opening="вступ",
                notes={"seo.h1.missing": "справжня", "seo.made_up": "вигадана"},
            )

    import app.report.owner_enrich as module

    original = module.get_analyst
    module.get_analyst = lambda *a, **k: Inventive()  # type: ignore[assignment]
    try:
        extras = await build_owner_extras(_report())
    finally:
        module.get_analyst = original

    assert "seo.made_up" not in extras.notes


def test_advice_from_the_model_never_reaches_the_owner():
    """Промпт її забороняє, але промпт — не гарантія. Порада «що зробити»
    у цьому документі віддає безкоштовно те, що продається."""
    from app.report.owner_enrich import looks_like_advice

    assert looks_like_advice("Додайте телефон у підвал сайту.")
    assert looks_like_advice("Проблема очевидна. Поставте ціни на видному місці.")
    assert looks_like_advice("Publish your prices.")
    # Опис проблеми з тим самим словом усередині — не порада.
    assert not looks_like_advice("Покупець не знає, куди додайте його б не привело меню.")
    assert not looks_like_advice("Ваш магазин не показує цін, і покупець іде далі.")


@pytest.mark.asyncio
async def test_a_note_that_slipped_into_advice_is_dropped_not_shown():
    import app.report.owner_enrich as module
    from app.audit.llm.base import OwnerNotes

    report = _report()
    good, bad = report.recommendations[0].msg, report.recommendations[1].msg

    class Chatty:
        name = "fake"

        async def owner_notes(self, business, band, problems):
            return OwnerNotes(
                analysed=True,
                opening="Ваш магазин втрачає покупців.",
                notes={good: "Покупець не бачить цін.", bad: "Додайте ціни на сторінку."},
            )

    original = module.get_analyst
    module.get_analyst = lambda *a, **k: Chatty()  # type: ignore[assignment]
    try:
        extras = await module.build_owner_extras(report)
    finally:
        module.get_analyst = original

    assert good in extras.notes
    assert bad not in extras.notes, "порада протекла у звіт для власника"
    assert extras.opening == "Ваш магазин втрачає покупців."
