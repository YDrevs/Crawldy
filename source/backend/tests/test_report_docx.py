"""Word-звіт: структура документа, дві мови, повнота каркасних рядків."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.audit.enums import AuditStatus, DegradationKind, Module, Severity
from app.audit.i18n import Locale
from app.audit.types import Degradation, Issue
from app.report.builder import build_report
from app.report.docx import filename_for, render_docx
from app.report.i18n import STRINGS
from app.report.i18n import t as _t
from app.report.types import ReportSource


def t(key: str, locale: Locale = Locale.UK, **params: object) -> str:
    """Рамкові рядки українською, бо українські й документи в цих тестах.

    Дефолт каталогу — англійська: вона канонічна. Тут потрібна саме та
    мова, якою побудований документ, інакше тест порівнює текст із
    перекладом іншої мови й падає з незрозумілим повідомленням.
    """
    return _t(key, locale, **params)

EVIDENCE = {
    "crawl": {"pages_ok": 5, "pages_failed": 1, "max_depth_reached": 2},
    "http": {"entry_url": "https://example.com/", "origin": "https://example.com",
             "is_https": True, "hsts": False},
    "robots_txt": {"exists": True, "status_code": 200, "blocked_ai_agents": ["GPTBot"]},
    "sitemap_xml": {"exists": False, "status_code": 404, "url_count": 0},
    "llms_txt": {"exists": False, "status_code": 404},
    "structured_data": {"pages_with_jsonld": 2, "pages_total": 5, "types": {"WebSite": 2}},
}


def _issue(msg, code, severity=Severity.MEDIUM, module=Module.SEO, page=None, params=None):
    return Issue(
        code=code,
        module=module,
        severity=severity,
        title="t",
        weight=2,
        page_url=page,
        msg=msg,
        params=params or {},
    )


def _source(**overrides) -> ReportSource:
    base: dict = dict(
        url="https://example.com",
        entry_url="https://example.com/",
        origin="https://example.com",
        overall=53,
        scores={"seo": 38, "ai_ready": 52, "llm_visibility": 76},
        issues=[
            _issue("seo.canonical.missing", "seo.canonical", page=f"https://example.com/{n}")
            for n in range(8)
        ]
        + [_issue("seo.sitemap.missing", "seo.sitemap", Severity.HIGH)],
        degraded=[
            Degradation(
                "performance",
                DegradationKind.PROVIDER_UNAVAILABLE,
                "429",
                [Module.PERFORMANCE],
            )
        ],
        skipped_modules={Module.PERFORMANCE},
        evidence=EVIDENCE,
        pages_crawled=5,
        status=AuditStatus.PARTIAL,
    )
    base.update(overrides)
    return ReportSource(**base)


def _document(locale=Locale.UK):
    report = build_report(_source(), locale)
    return Document(io.BytesIO(render_docx(report, locale)))


def _paragraphs(document) -> list[str]:
    return [p.text for p in document.paragraphs if p.text.strip()]


def _tables(document) -> list[list[list[str]]]:
    return [[[c.text for c in row.cells] for row in table.rows] for table in document.tables]


def _headings(document, level: str) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name == level and p.text.strip()]


def test_docx_is_a_valid_openxml_package():
    data = render_docx(build_report(_source(), Locale.UK))
    assert data[:2] == b"PK", "docx — це zip-контейнер"
    assert len(data) > 10_000
    Document(io.BytesIO(data))  # не кидає — файл читається


@pytest.mark.parametrize("locale", list(Locale))
def test_document_splits_into_two_parts(locale: Locale):
    headings = _headings(_document(locale), "Heading 1")

    assert headings == [t("part1_title", locale), t("part2_title", locale)]


@pytest.mark.parametrize("locale", list(Locale))
def test_each_part_has_its_own_sections(locale: Locale):
    document = _document(locale)
    sections = _headings(document, "Heading 2")

    for key in ("section_scores", "section_priorities", "section_missing"):
        assert t(key, locale) in sections, f"частина 1 без «{key}»"
    for key in (
        "section_scores_technical",
        "section_missing_technical",
        "section_details",
        "section_evidence",
    ):
        assert t(key, locale) in sections, f"частина 2 без «{key}»"


def test_section_headings_are_not_reused_across_parts():
    """Однакова назва двічі ламає навігацію по документу."""
    sections = _headings(_document(), "Heading 2")
    assert len(sections) == len(set(sections)), f"дублі заголовків: {sections}"


def test_plain_part_comes_before_the_technical_one():
    order = _headings(_document(), "Heading 2")

    assert order.index(t("section_priorities")) < order.index(t("section_details"))
    assert order.index(t("section_scores")) < order.index(t("section_scores_technical"))


def test_plain_part_explains_each_area_in_words():
    """Нетехнічна таблиця пояснює, що взагалі означає напрям."""
    areas_table = _tables(_document())[0]
    header, *rows = areas_table

    assert header[1] == t("col_area_meaning")
    seo_row = next(r for r in rows if r[0] == t("module_seo"))
    assert seo_row[1] == t("area_seo")
    assert "пошуков" in seo_row[1].lower(), "пояснення без жаргону"


def test_plain_part_leads_with_the_business_consequence():
    """Спершу «що це коштує», потім «що зробити» — читає той, хто вирішує."""
    paragraphs = _paragraphs(_document())
    priorities_at = paragraphs.index(t("section_priorities"))
    tail = paragraphs[priorities_at:]

    assert any(p.startswith("1. ") for p in tail)
    first = next(p for p in tail if p.startswith("1. "))
    # Раніше тут стояли слова «пошук» і «сторінк» — вони жили в тодішньому тексті
    # про карту сайту, і перевірка трималася на них, а не на інваріанті. Щойно
    # затверджений наслідок перестав містити ці слова, тест упав, хоч документ і далі
    # веде наслідком. Тепер порівнюється те, що тест і стверджує.
    report = build_report(_source(), Locale.UK)
    assert first[len("1. "):] == next(
        r for r in report.recommendations if r.consequence
    ).consequence, f"перший пріоритет має вести наслідком, а не дією: {first}"
    assert any(p.startswith(f"{t('what_to_do')}: ") for p in tail)


def test_plain_part_hides_urls_and_evidence():
    """Частина 1 не має технічних деталей — вони починаються з частини 2."""
    paragraphs = _paragraphs(_document())
    part2_at = paragraphs.index(t("part2_title"))
    plain = paragraphs[:part2_at]

    assert not [p for p in plain if "https://example.com/" in p], "у частині 1 не має бути URL"
    assert t("evidence") not in " ".join(plain)


def test_scores_table_marks_skipped_module_as_not_measured():
    table = _tables(_document())[0]
    speed_row = next(r for r in table if r[0] == t("module_performance"))

    assert speed_row[2] == t("not_measured")
    assert speed_row[3] == "—", "смуги в пропущеного напряму немає"


def test_technical_table_reports_data_completeness():
    partial_source = _source(
        skipped_modules=set(),
        partial_modules={Module.PERFORMANCE},
        scores={"performance": 60, "seo": 38, "ai_ready": 52, "llm_visibility": 76},
    )
    document = Document(io.BytesIO(render_docx(build_report(partial_source, Locale.UK))))
    technical = _tables(document)[1]
    speed_row = next(r for r in technical if r[0] == t("module_performance"))

    assert speed_row[1] == "60"
    assert speed_row[5] == t("completeness_partial")


def test_priorities_carry_gain_and_effort():
    text = "\n".join(_paragraphs(_document()))

    assert t("gives") in text
    assert t("effort") in text
    assert t("impact_note") in text, "клієнт не має сумувати ефекти"


def test_technical_part_still_shows_scope_and_evidence():
    paragraphs = _paragraphs(_document())
    part2_at = paragraphs.index(t("part2_title"))
    technical = " ".join(paragraphs[part2_at:])

    assert t("site_wide") in technical
    assert t("evidence") in technical
    assert "https://example.com/0" in technical


def test_full_breakdown_covers_every_module_even_when_clean():
    document = _document()
    module_headings = _headings(document, "Heading 3")

    assert module_headings == [t(f"module_{m.value}") for m in Module]
    assert t("module_not_scored") in _paragraphs(document)


def test_page_examples_are_truncated_with_a_counter():
    """8 сторінок, показуємо 5 і кажемо, скільки лишилось."""
    document = _document()
    paragraphs = _paragraphs(document)

    assert "https://example.com/0" in paragraphs
    assert t("and_more", count=3) in paragraphs


def test_evidence_appendix_reports_what_was_seen():
    rows = {r[0]: r[1] for r in _tables(_document())[-1]}

    assert rows[t("ev_pages_ok")] == "5"
    assert rows[t("ev_https")] == t("yes")
    assert rows[t("ev_hsts")] == t("no")
    assert rows[t("ev_sitemap")] == f"{t('absent')} (HTTP 404)"
    assert rows[t("ev_blocked_agents")] == "GPTBot"


def test_language_changes_prose_but_not_numbers():
    uk, en = _document(Locale.UK), _document(Locale.EN)

    assert "Аудит сайту" in _paragraphs(uk)[0]
    assert "Site audit" in _paragraphs(en)[0]
    seo_uk = next(r for r in _tables(uk)[0] if r[0] == t("module_seo", Locale.UK))
    seo_en = next(r for r in _tables(en)[0] if r[0] == t("module_seo", Locale.EN))
    assert seo_uk[1] != seo_en[1], "пояснення напряму перекладається"
    assert seo_uk[2] == seo_en[2] == "38", "оцінка від мови не залежить"


def test_filename_encodes_host_and_language():
    report = build_report(_source(), Locale.UK)
    assert filename_for(report, Locale.UK) == "audit-example.com-technical-uk.docx"
    assert filename_for(report, Locale.EN) == "audit-example.com-technical-en.docx"


# ── каркасні рядки ─────────────────────────────────────────────────────────


def test_report_string_catalogs_have_identical_keys():
    uk, en = STRINGS[Locale.UK], STRINGS[Locale.EN]
    assert set(uk) == set(en), f"розбіжність: {set(uk) ^ set(en)}"


def test_every_band_effort_and_severity_has_a_label():
    for locale in Locale:
        for band in ("critical", "poor", "medium", "good", "unknown"):
            assert t(f"band_{band}", locale) != f"band_{band}"
        for effort in ("low", "medium", "high"):
            assert t(f"effort_{effort}", locale) != f"effort_{effort}"
        for severity in Severity:
            assert t(f"severity_{severity.value}", locale) != f"severity_{severity.value}"
        for module in Module:
            assert t(f"module_{module.value}", locale) != f"module_{module.value}"


def test_urgent_is_reserved_for_critical_findings():
    """Слово «терміново» знецінюється, якщо стоїть у кожному звіті.

    `high` мав той самий підпис, що й `critical`, а HIGH-знахідки має
    майже кожен сайт: seo.h1 спрацьовує на 21 із 29, tech.analytics на
    19. На ka-dent.com.ua пʼять рядків «ТЕРМІНОВО» стояли при нулі
    критичних знахідок.
    """
    from app.report.docx_owner import SEVERITY_BAND

    assert SEVERITY_BAND["critical"] != SEVERITY_BAND["high"]
    assert SEVERITY_BAND["critical"] == "critical"


def test_every_severity_has_its_own_words_for_the_owner():
    """Два однакові підписи на різні рівні — це підказка нізащо."""
    from app.audit.enums import Severity
    from app.report.docx_owner import SEVERITY_BAND
    from app.report.i18n import STRINGS

    for locale, strings in STRINGS.items():
        labels = [
            strings[f"owner_gravity_{SEVERITY_BAND[s.value]}"] for s in Severity
        ]
        assert len(set(labels)) == len(labels), f"{locale}: {labels}"
