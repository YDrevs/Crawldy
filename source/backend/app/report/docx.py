"""Рендер звіту у Word.

python-docx, без системних залежностей — свідомо замість PDF через
headless Chrome або WeasyPrint з GTK.

Документ ділиться на дві частини під двох різних читачів.

Частина 1 — для того, хто ухвалює рішення: вердикт прозою, чотири напрями
простою мовою, пріоритети, подані через наслідок для бізнесу. Ні URL,
ні доказів, ні ваг — усе, що вимагає технічного контексту, лишається далі.

Частина 2 — для того, хто виправляє: оцінки з вагами й повнотою даних,
повний розбір із доказами та адресами сторінок, додаток зі спостереженнями.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt

from app.audit.enums import Module
from app.audit.i18n import Locale
from app.report.docx_brand import (
    MUTED,
    add_cover,
    add_footer,
    mono_run,
    set_column_widths,
    severity_run,
    style_document,
    style_table,
)
from app.report.i18n import locale_of, t
from app.report.types import (
    MODULE_FIELD,
    Recommendation,
    ReportDocument,
)

MAX_EXAMPLES_SHOWN = 5


def render_docx(report: ReportDocument, locale: Locale | None = None) -> bytes:
    """Один документ, дві частини під двох різних читачів.

    Частина 1 — для того, хто ухвалює рішення: вердикт прозою, чотири
    напрями простою мовою, пріоритети через наслідок. Жодних URL і доказів.
    Частина 2 — для того, хто виправляє: усе решта.

    Мова за замовчуванням береться **з документа**, а не з дефолту
    продукту. Різниця не теоретична: документ уже зібраний якоюсь мовою,
    і другий незалежний аргумент означав два джерела правди про одну
    річ. Вони й розійшлися — рамкові заголовки виходили англійськими на
    українському звіті, а зміст лишався українським.
    """
    locale = locale or locale_of(report)
    document = Document()
    style_document(document)
    add_cover(document, report, locale, kind="kind_technical")
    add_footer(document, report)

    _part_heading(document, report, locale, "part1_title", "part1_subtitle")
    _verdict(document, report, locale)
    _areas_table(document, report, locale)
    _priorities(document, report, locale)
    _missing_plain(document, report, locale)

    document.add_page_break()
    _part_heading(document, report, locale, "part2_title", "part2_subtitle")
    _scores_table(document, report, locale)
    _accessibility(document, report, locale)
    _missing_data(document, report, locale)
    _details(document, report, locale)
    _passed(document, report, locale)
    _evidence(document, report, locale)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _part_heading(
    document: Any, report: ReportDocument, locale: Locale, title_key: str, subtitle_key: str
) -> None:
    document.add_heading(t(title_key, locale), level=1)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(t(subtitle_key, locale))
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


# ── частина 1: для нетехнічного читача ─────────────────────────────────────


def _verdict(document: Any, report: ReportDocument, locale: Locale) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(t(f"verdict_{report.overall_band}", locale))
    run.font.size = Pt(11)


def _areas_table(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Без ваг, відсотків і кількості знахідок — лише що це і як справи."""
    document.add_heading(t("section_scores", locale), level=2)

    table = document.add_table(rows=1, cols=4)
    headers = ("col_module", "col_area_meaning", "col_score", "col_band")
    for cell, key in zip(table.rows[0].cells, headers, strict=True):
        cell.text = t(key, locale)

    for entry in report.scores:
        row = table.add_row().cells
        row[0].text = t(f"module_{entry.module.value}", locale)
        row[1].text = t(f"area_{entry.module.value}", locale)
        row[2].text = t("not_measured", locale) if entry.excluded else str(entry.score)
        verdict = "—" if entry.excluded else t(f"band_{entry.band}", locale)
        if entry.partial:
            verdict = f"{verdict} ({t('completeness_partial', locale)})"
        row[3].text = verdict
    style_table(table)
    set_column_widths(table, (3.6, 7.6, 2.2, 3.1))


def _accessibility(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Знімок доступності: одна сторінка, лічильники, кожен провал.

    Бала 0..100 тут немає навмисно, і «100/100» немає теж: чиста
    сторінка описується кількістю перевірок, які на ній запустились.
    Показуємо КОЖНУ провалену перевірку Lighthouse, включно з тими, що
    відсіяні з входу чеків, — фільтр дублів вирішує, за що штрафувати,
    а не що клієнтові знати.
    """
    snapshot = report.accessibility
    document.add_heading(t("section_a11y_snapshot", locale), level=2)
    if snapshot is None or not snapshot.available:
        document.add_paragraph(t("a11y_unavailable", locale))
        _weight_note(document, snapshot, locale)
        return

    document.add_paragraph(t("a11y_scope", locale, url=snapshot.measured_url or "—"))
    if snapshot.failed:
        document.add_paragraph(
            t(
                "a11y_counts",
                locale,
                applicable=snapshot.applicable,
                passed=snapshot.passed,
                failed=snapshot.failed,
            )
        )
        table = document.add_table(rows=1, cols=3)
        headers = ("col_a11y_check", "col_a11y_impact", "col_a11y_also")
        for cell, key in zip(table.rows[0].cells, headers, strict=True):
            cell.text = t(key, locale)
        for finding in snapshot.findings:
            row = table.add_row().cells
            row[0].text = finding.title or finding.audit_id
            row[1].text = finding.impact or "—"
            # Посилання ставимо лише там, де відповідна знахідка справді
            # є в цьому аудиті. Мапа правил сама по собі покриття не
            # доводить, і обіцяти його «за мапою» означало б відправити
            # клієнта шукати те, чого в звіті немає.
            row[2].text = (
                t(
                    "a11y_also_reported",
                    locale,
                    module=t(f"module_{finding.also_reported_module}", locale),
                )
                if finding.also_reported_as
                else "—"
            )
        style_table(table)
        set_column_widths(table, (7.4, 3.0, 6.1))
    else:
        document.add_paragraph(t("a11y_clean", locale, applicable=snapshot.applicable))

    if snapshot.lighthouse_score is not None:
        reference = document.add_paragraph(
            t("a11y_lighthouse_ref", locale, score=f"{snapshot.lighthouse_score:.0f}")
        )
        reference.runs[0].italic = True
    caveat = document.add_paragraph(t("a11y_caveat", locale))
    caveat.runs[0].italic = True
    _weight_note(document, snapshot, locale)


def _weight_note(document: Any, snapshot: Any, locale: Locale) -> None:
    """Чи впливає цей знімок на загальний бал — і наскільки."""
    if snapshot is not None and snapshot.counts_toward_overall:
        text = t(
            "a11y_still_weighted",
            locale,
            percent=f"{snapshot.overall_weight or 0.0:.0%}",
        )
    else:
        text = t("a11y_not_weighted", locale)
    note = document.add_paragraph(text)
    note.runs[0].italic = True


def _priorities(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Веде наслідок, а не дія: спершу «що це коштує», потім «що зробити»."""
    document.add_heading(t("section_priorities", locale), level=2)
    if not report.next_actions:
        document.add_paragraph(t("no_priorities", locale))
        return

    note = document.add_paragraph()
    note.add_run(t("impact_note", locale)).italic = True
    note.runs[0].font.size = Pt(9)

    for action in report.next_actions:
        headline = document.add_paragraph()
        headline.add_run(f"{action.rank}. ").bold = True
        headline.add_run(action.why).bold = True

        todo = document.add_paragraph()
        todo.paragraph_format.left_indent = Pt(18)
        todo.add_run(f"{t('what_to_do', locale)}: ").bold = True
        todo.add_run(action.action)

        facts = document.add_paragraph()
        facts.paragraph_format.left_indent = Pt(18)
        facts.add_run(
            f"{t('gives', locale)}: {action.expected_gain}  ·  "
            f"{t('effort', locale)}: {t(f'effort_{action.effort.value}', locale)}"
        ).font.size = Pt(9)


def _missing_plain(document: Any, report: ReportDocument, locale: Locale) -> None:
    if not report.missing_data:
        return
    document.add_heading(t("section_missing", locale), level=2)
    document.add_paragraph(t("part1_missing_intro", locale))
    for entry in report.missing_data:
        affected = ", ".join(
            t(f"module_{m.value}", locale)
            for m in Module
            if MODULE_FIELD[m] in entry.affects
        )
        document.add_paragraph(affected or entry.scope, style="List Bullet")


def filename_for(report: ReportDocument, locale: Locale) -> str:
    host = report.target.get("origin", "").replace("https://", "").replace("http://", "")
    safe = "".join(c if c.isalnum() or c in "-._" else "-" for c in host) or "audit"
    # Вид у назві файлу: чотири документи одного сайту не мають
    # перезаписувати один одного в теці завантажень.
    return f"audit-{safe}-technical-{locale.value}.docx"


def _scores_table(document: Any, report: ReportDocument, locale: Locale) -> None:
    document.add_heading(t("section_scores_technical", locale), level=2)

    table = document.add_table(rows=1, cols=6)
    headers = (
        "col_module",
        "col_score",
        "col_band",
        "col_issues",
        "col_weight",
        "col_completeness",
    )
    for cell, key in zip(table.rows[0].cells, headers, strict=True):
        cell.text = t(key, locale)

    for entry in report.scores:
        row = table.add_row().cells
        row[0].text = t(f"module_{entry.module.value}", locale)
        row[1].text = t("not_measured", locale) if entry.excluded else str(entry.score)
        row[2].text = "—" if entry.excluded else t(f"band_{entry.band}", locale)
        row[3].text = str(entry.issues_count)
        row[4].text = "—" if entry.excluded else f"{entry.effective_weight:.0%}"
        if entry.excluded:
            row[5].text = "—"
        else:
            row[5].text = t(
                "completeness_partial" if entry.partial else "completeness_full", locale
            )
    style_table(table)
    set_column_widths(table, (3.6, 1.8, 2.6, 2.2, 2.4, 3.9))
    _weight_gap(document, report, locale)


def _weight_gap(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Куди подівся десяток відсотків ваги.

    Доступність пішла з таблиці, а її вклад у загальний бал лишився —
    тож видимі ваги сходяться до 90%, а не до 100%. Мовчати про це не
    можна двічі: читач або вирішить, що звіт зламаний, або — значно
    гірше — що знімок доступності на його бал не впливає взагалі.
    """
    snapshot = report.accessibility
    if snapshot is None or not snapshot.counts_toward_overall:
        return
    hidden = snapshot.overall_weight or 0.0
    shown = sum(entry.effective_weight for entry in report.scores)
    note = document.add_paragraph(
        t("scores_weight_gap", locale, shown=f"{shown:.0%}", hidden=f"{hidden:.0%}")
    )
    note.runs[0].italic = True


# ── прогалини ──────────────────────────────────────────────────────────────


def _passed(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Пройдені перевірки: компактно, після проблем і дій.

    Word не вміє згортати, тож секція просто друга за порядком і
    стримана за виглядом: згруповані по модулях однорядкові пункти без
    зелених плакатів. Нуль пройдених — секції немає взагалі.
    """
    if not report.passed_checks:
        return
    document.add_heading(
        t("section_passed", locale, count=len(report.passed_checks)), level=2
    )
    intro = document.add_paragraph(t("passed_intro", locale))
    intro.runs[0].italic = True
    by_module: dict[str, list[str]] = {}
    for check in report.passed_checks:
        title = t(f"passed.{check.key}", locale, **check.facts)
        by_module.setdefault(check.module.value, []).append(title)
    for module_value, titles in by_module.items():
        line = document.add_paragraph(style="List Bullet")
        line.add_run(t(f"module_{module_value}", locale) + ": ").bold = True
        line.add_run("; ".join(titles))


def _missing_data(document: Any, report: ReportDocument, locale: Locale) -> None:
    """Технічна версія: зі scope, причиною і тим, на який скор це вплинуло."""
    if not report.missing_data:
        return
    document.add_heading(t("section_missing_technical", locale), level=2)
    for entry in report.missing_data:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{entry.scope} ({entry.kind}). ").bold = True
        paragraph.add_run(f"{t('reason', locale)}: {entry.reason}. ")
        if entry.affects:
            paragraph.add_run(f"{t('affects', locale)}: {', '.join(entry.affects)}. ")
        paragraph.add_run(entry.note).italic = True


# ── повний розбір ──────────────────────────────────────────────────────────


def _details(document: Any, report: ReportDocument, locale: Locale) -> None:
    document.add_heading(t("section_details", locale), level=2)
    by_module: dict[Module, list[Recommendation]] = {m: [] for m in MODULE_FIELD}
    for recommendation in report.recommendations:
        by_module[recommendation.module].append(recommendation)

    excluded = {s.module for s in report.scores if s.excluded}
    for module, recommendations in by_module.items():
        document.add_heading(t(f"module_{module.value}", locale), level=3)
        if module in excluded:
            document.add_paragraph(t("module_not_scored", locale))
            continue
        if not recommendations:
            document.add_paragraph(t("no_issues", locale))
            continue
        for recommendation in recommendations:
            _recommendation(document, recommendation, locale)


def _recommendation(document: Any, rec: Recommendation, locale: Locale) -> None:
    document.add_heading(rec.title, level=4)

    effect = f"+{rec.impact_points}"
    if rec.floored:
        effect = f"{effect} ({t('impact_floored', locale)})"

    facts = document.add_paragraph()
    # Заголовок не має лишатися сиротою над розривом сторінки без
    # свого рядка фактів.
    facts.paragraph_format.keep_with_next = True
    pages = t("site_wide", locale) if rec.scope == "site" else str(rec.affected_pages)
    severity_run(
        facts, f"{t('severity', locale)}: {t(f'severity_{rec.severity.value}', locale)}",
        rec.severity.value,
    )
    facts.add_run(
        f"  ·  {t('effect', locale)}: {effect}  ·  "
        f"{t('effort', locale)}: {t(f'effort_{rec.effort.value}', locale)}  ·  "
        f"{t('pages_affected', locale)}: {pages}"
    ).font.size = Pt(9)

    if rec.problem:
        _labelled(document, t("problem", locale), rec.problem)
    if rec.action:
        _labelled(document, t("what_to_do", locale), rec.action)
    if rec.evidence:
        # Доказ — технічне значення: моноширинним і дрібніше, щоб
        # відрізнявся від тексту рекомендації з першого погляду.
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{t('evidence', locale)}: ").bold = True
        mono_run(paragraph, _format_evidence(rec.evidence))

    if rec.affected_examples:
        document.add_paragraph(f"{t('examples', locale)}:").runs[0].bold = True
        for url in rec.affected_examples[:MAX_EXAMPLES_SHOWN]:
            bullet = document.add_paragraph(style="List Bullet")
            mono_run(bullet, url)
        hidden = rec.affected_pages - MAX_EXAMPLES_SHOWN
        if hidden > 0:
            document.add_paragraph(t("and_more", locale, count=hidden), style="List Bullet")


def _labelled(document: Any, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text)


# Глибина, до якої розгортаємо вкладений доказ. Далі — лише розмір.
# Доказ атрибуції — це список цілей, у кожній список прикладів; без межі
# у Word їхав сирий `repr` словника з лапками й `False`, тобто рядок,
# який клієнт прочитати не може.
_EVIDENCE_DEPTH = 2
# Скільки елементів списку показуємо. Решту згортаємо в «+N», а не
# викидаємо мовчки: урізаний перелік не має вдавати повний.
_EVIDENCE_ITEMS = 3


def _short(value: Any, depth: int = 0) -> str:
    """Вкладений доказ у рядок, придатний для читання очима.

    Адреси й текст посилань не перекладаються — це дані. Перекладаються
    підписи навколо них, і вони приходять із каталогу окремо.
    """
    if isinstance(value, dict):
        if not value:
            return "—"
        if depth >= _EVIDENCE_DEPTH:
            return f"{{{len(value)}}}"
        return " ".join(f"{k}={_short(v, depth + 1)}" for k, v in list(value.items())[:6])
    if isinstance(value, list):
        if not value:
            return "—"
        if depth >= _EVIDENCE_DEPTH:
            return f"[{len(value)}]"
        shown = [_short(item, depth + 1) for item in value[:_EVIDENCE_ITEMS]]
        rest = len(value) - len(shown)
        return " | ".join(shown) + (f" (+{rest})" if rest > 0 else "")
    if value is None:
        return "—"
    if isinstance(value, bool):
        # Мовно нейтрально: це дані, а не підпис.
        return "true" if value else "false"
    return str(value)


def _format_evidence(evidence: dict[str, Any]) -> str:
    return "; ".join(f"{key} = {_short(value)}" for key, value in evidence.items())


# ── додаток ────────────────────────────────────────────────────────────────


def _evidence(document: Any, report: ReportDocument, locale: Locale) -> None:
    if not report.evidence:
        return
    document.add_heading(t("section_evidence", locale), level=2)

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = t("col_parameter", locale)
    table.rows[0].cells[1].text = t("col_value", locale)

    for label, value in _evidence_rows(report.evidence, locale):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
    style_table(table)
    set_column_widths(table, (8.2, 8.3))


def _evidence_rows(evidence: dict[str, Any], locale: Locale) -> list[tuple[str, str]]:
    crawl = evidence.get("crawl", {})
    http = evidence.get("http", {})
    robots = evidence.get("robots_txt", {})
    sitemap = evidence.get("sitemap_xml", {})
    llms = evidence.get("llms_txt", {})
    structured = evidence.get("structured_data", {})

    def flag(value: object) -> str:
        return t("yes", locale) if value else t("no", locale)

    def presence(info: dict[str, Any]) -> str:
        status = info.get("status_code")
        mark = t("present", locale) if info.get("exists") else t("absent", locale)
        return f"{mark} (HTTP {status})" if status else mark

    types = structured.get("types") or {}
    blocked = robots.get("blocked_ai_agents") or []
    return [
        (t("ev_pages_ok", locale), str(crawl.get("pages_ok", 0))),
        (t("ev_pages_failed", locale), str(crawl.get("pages_failed", 0))),
        (t("ev_max_depth", locale), str(crawl.get("max_depth_reached", 0))),
        (t("ev_https", locale), flag(http.get("is_https"))),
        (t("ev_hsts", locale), flag(http.get("hsts"))),
        (t("ev_robots", locale), presence(robots)),
        (t("ev_sitemap", locale), presence(sitemap)),
        (t("ev_url_count", locale), str(sitemap.get("url_count", 0))),
        (t("ev_llms", locale), presence(llms)),
        (
            t("ev_jsonld", locale),
            f"{structured.get('pages_with_jsonld', 0)} / {structured.get('pages_total', 0)}",
        ),
        (t("ev_schema_types", locale), ", ".join(types) or t("none", locale)),
        (t("ev_blocked_agents", locale), ", ".join(blocked) or t("none", locale)),
    ]


__all__ = ["filename_for", "render_docx"]
