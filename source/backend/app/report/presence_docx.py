"""Word-експорт Actual AI Presence: факти спостережень, без обіцянок.

Окремий документ, НЕ частина аудиторського звіту: присутність — окрема
продуктова секція з окремою методологією, і склеювання її з
детерміністичним аудитом розмивало б обидва. Вхід — прості словники
(жодного ORM і FastAPI — та сама межа, що в усього report/).

Мови — всі шість повних локалей звітів (en/uk/de/es/fr/ru): експорт —
клієнтський документ, і він мусить бути мовою клієнта, як і решта
звітності. Кабінет лишається en/uk за рішенням продукту — це інша
поверхня. Формулювання рахує «x із N перевірених спостережень» і ніколи
не обіцяє ні відсотка видимості, ні поведінки споживчих застосунків.
Назви провайдерів — офіційні торгові назви, вони не перекладаються.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from docx import Document
from docx.shared import Pt

from app.audit.i18n import DEFAULT_LOCALE, parse
from app.report.docx_brand import MUTED, add_footer_for, style_document
from app.report.i18n import t as label

# Офіційні назви провайдерів — номінативне використання торгових назв;
# логотипи не використовуються (правила брендів вимагають окремого
# дозволу, і текстова назва — єдине завжди дозволене подання).
PROVIDER_NAMES = {
    "openai": "ChatGPT / OpenAI API",
    "anthropic": "Claude / Anthropic API",
    "perplexity": "Perplexity API",
    "google": "Gemini / Google API",
}

_STRINGS: dict[str, dict[str, str]] = {
    "en": {
        "title": "Actual AI Presence",
        "for": "Website",
        "brand": "Brand",
        "generated": "Generated",
        "method": (
            "Method: controlled checks through the providers' public APIs "
            "with web search enabled. Results are stochastic; API "
            "observations may differ from the consumer ChatGPT, Claude, "
            "Perplexity or Gemini applications. This report never claims "
            "rankings and is independent of the providers; no affiliation "
            "or endorsement is implied."
        ),
        "modes_natural": "Natural answers (the model decides whether to search)",
        "modes_explicit": "Explicitly web-grounded answers",
        "observations": "valid observations",
        "mentioned": "brand mentioned",
        "cited": "site cited as a source",
        "unavailable": "observations unavailable",
        "coverage": "Provider coverage",
        "competitors": "Most-cited domains across observations",
        "evidence": "Observation evidence (excerpts)",
        "deliverable_yes": (
            "This run met the coverage threshold: one AI Presence check "
            "was used."
        ),
        "deliverable_no": (
            "This run did not meet the coverage threshold: no AI Presence "
            "check was charged."
        ),
        "searched": "searched the web",
        "no_search": "answered without searching",
        "status_prefix": "status",
        "marks": (
            "Provider names and marks belong to their respective owners. "
            "Crawldy is an independent service and is not affiliated with "
            "or endorsed by the providers shown."
        ),
        "no_observations": (
            "No valid observations were completed for this provider; its "
            "results are reported as unavailable, not as zero presence."
        ),
    },
    "uk": {
        "title": "Actual AI Presence",
        "for": "Сайт",
        "brand": "Бренд",
        "generated": "Згенеровано",
        "method": (
            "Метод: контрольовані перевірки через публічні API провайдерів "
            "з увімкненим веб-пошуком. Результати стохастичні; "
            "API-спостереження можуть відрізнятися від споживчих "
            "застосунків ChatGPT, Claude, Perplexity чи Gemini. Звіт не "
            "стверджує позицій у видачі та незалежний від провайдерів; "
            "жодної афіліації чи схвалення не мається на увазі."
        ),
        "modes_natural": "Природні відповіді (модель сама вирішує, чи шукати)",
        "modes_explicit": "Явно заземлені у веб відповіді",
        "observations": "перевірених спостережень",
        "mentioned": "бренд згадано",
        "cited": "сайт процитовано як джерело",
        "unavailable": "спостережень недоступно",
        "coverage": "Покриття провайдерів",
        "competitors": "Найцитованіші домени у спостереженнях",
        "evidence": "Докази спостережень (уривки)",
        "deliverable_yes": (
            "Прогін досяг порога покриття: використано одну перевірку "
            "AI-присутності."
        ),
        "deliverable_no": (
            "Прогін не досяг порога покриття: перевірку AI-присутності "
            "не списано."
        ),
        "searched": "шукав у вебі",
        "no_search": "відповів без пошуку",
        "status_prefix": "статус",
        "marks": (
            "Назви й знаки провайдерів належать їхнім власникам. Crawldy — "
            "незалежний сервіс, не афілійований із показаними провайдерами "
            "і не схвалений ними."
        ),
        "no_observations": (
            "Для цього провайдера не завершено жодного валідного "
            "спостереження; його результат подається як недоступний, а не "
            "як нульова присутність."
        ),
    },
    "de": {
        "title": "Actual AI Presence",
        "for": "Website",
        "brand": "Marke",
        "generated": "Erstellt",
        "method": (
            "Methode: kontrollierte Abfragen über die öffentlichen APIs der "
            "Anbieter mit aktivierter Websuche. Die Ergebnisse sind "
            "stochastisch; API-Beobachtungen können von den Verbraucher-Apps "
            "ChatGPT, Claude, Perplexity oder Gemini abweichen. Dieser "
            "Bericht behauptet keine Platzierungen und ist von den Anbietern "
            "unabhängig; es wird keine Verbindung oder Befürwortung "
            "unterstellt."
        ),
        "modes_natural": "Natürliche Antworten (das Modell entscheidet, ob es sucht)",
        "modes_explicit": "Ausdrücklich webgestützte Antworten",
        "observations": "gültige Beobachtungen",
        "mentioned": "Marke erwähnt",
        "cited": "Website als Quelle zitiert",
        "unavailable": "Beobachtungen nicht verfügbar",
        "coverage": "Anbieterabdeckung",
        "competitors": "Meistzitierte Domains in den Beobachtungen",
        "evidence": "Beobachtungsbelege (Auszüge)",
        "deliverable_yes": (
            "Dieser Lauf erreichte die Abdeckungsschwelle: eine "
            "AI-Presence-Prüfung wurde verbraucht."
        ),
        "deliverable_no": (
            "Dieser Lauf erreichte die Abdeckungsschwelle nicht: es wurde "
            "keine AI-Presence-Prüfung berechnet."
        ),
        "searched": "hat im Web gesucht",
        "no_search": "antwortete ohne Suche",
        "status_prefix": "Status",
        "marks": (
            "Anbieternamen und -marken gehören ihren jeweiligen Inhabern. "
            "Crawldy ist ein unabhängiger Dienst und steht in keiner "
            "Verbindung zu den gezeigten Anbietern und wird von ihnen nicht "
            "befürwortet."
        ),
        "no_observations": (
            "Für diesen Anbieter wurde keine gültige Beobachtung "
            "abgeschlossen; sein Ergebnis wird als nicht verfügbar "
            "ausgewiesen, nicht als null Präsenz."
        ),
    },
    "es": {
        "title": "Actual AI Presence",
        "for": "Sitio web",
        "brand": "Marca",
        "generated": "Generado",
        "method": (
            "Método: comprobaciones controladas a través de las API públicas "
            "de los proveedores con búsqueda web activada. Los resultados son "
            "estocásticos; las observaciones por API pueden diferir de las "
            "aplicaciones de consumo ChatGPT, Claude, Perplexity o Gemini. "
            "Este informe no afirma posiciones y es independiente de los "
            "proveedores; no se implica afiliación ni respaldo."
        ),
        "modes_natural": "Respuestas naturales (el modelo decide si buscar)",
        "modes_explicit": "Respuestas explícitamente basadas en la web",
        "observations": "observaciones válidas",
        "mentioned": "marca mencionada",
        "cited": "sitio citado como fuente",
        "unavailable": "observaciones no disponibles",
        "coverage": "Cobertura por proveedor",
        "competitors": "Dominios más citados en las observaciones",
        "evidence": "Evidencia de observaciones (extractos)",
        "deliverable_yes": (
            "Esta ejecución alcanzó el umbral de cobertura: se consumió una "
            "comprobación de AI Presence."
        ),
        "deliverable_no": (
            "Esta ejecución no alcanzó el umbral de cobertura: no se cobró "
            "ninguna comprobación de AI Presence."
        ),
        "searched": "buscó en la web",
        "no_search": "respondió sin buscar",
        "status_prefix": "estado",
        "marks": (
            "Los nombres y las marcas de los proveedores pertenecen a sus "
            "respectivos propietarios. Crawldy es un servicio independiente y "
            "no está afiliado a los proveedores mostrados ni cuenta con su "
            "respaldo."
        ),
        "no_observations": (
            "No se completó ninguna observación válida para este proveedor; "
            "su resultado se presenta como no disponible, no como presencia "
            "cero."
        ),
    },
    "fr": {
        "title": "Actual AI Presence",
        "for": "Site web",
        "brand": "Marque",
        "generated": "Généré",
        "method": (
            "Méthode : vérifications contrôlées via les API publiques des "
            "fournisseurs avec recherche web activée. Les résultats sont "
            "stochastiques ; les observations par API peuvent différer des "
            "applications grand public ChatGPT, Claude, Perplexity ou "
            "Gemini. Ce rapport ne revendique aucun classement et est "
            "indépendant des fournisseurs ; aucune affiliation ni "
            "approbation n'est sous-entendue."
        ),
        "modes_natural": "Réponses naturelles (le modèle décide de chercher ou non)",
        "modes_explicit": "Réponses explicitement ancrées dans le web",
        "observations": "observations valides",
        "mentioned": "marque mentionnée",
        "cited": "site cité comme source",
        "unavailable": "observations indisponibles",
        "coverage": "Couverture par fournisseur",
        "competitors": "Domaines les plus cités dans les observations",
        "evidence": "Preuves d'observation (extraits)",
        "deliverable_yes": (
            "Cette exécution a atteint le seuil de couverture : une "
            "vérification AI Presence a été consommée."
        ),
        "deliverable_no": (
            "Cette exécution n'a pas atteint le seuil de couverture : "
            "aucune vérification AI Presence n'a été facturée."
        ),
        "searched": "a cherché sur le web",
        "no_search": "a répondu sans chercher",
        "status_prefix": "statut",
        "marks": (
            "Les noms et marques des fournisseurs appartiennent à leurs "
            "propriétaires respectifs. Crawldy est un service indépendant, "
            "sans affiliation avec les fournisseurs présentés et sans "
            "approbation de leur part."
        ),
        "no_observations": (
            "Aucune observation valide n'a été réalisée pour ce "
            "fournisseur ; son résultat est présenté comme indisponible, et "
            "non comme une présence nulle."
        ),
    },
    "ru": {
        "title": "Actual AI Presence",
        "for": "Сайт",
        "brand": "Бренд",
        "generated": "Сформировано",
        "method": (
            "Метод: контролируемые проверки через публичные API провайдеров "
            "с включённым веб-поиском. Результаты стохастичны; наблюдения "
            "через API могут отличаться от потребительских приложений "
            "ChatGPT, Claude, Perplexity или Gemini. Отчёт не утверждает "
            "позиций и независим от провайдеров; аффилиация или одобрение "
            "не подразумеваются."
        ),
        "modes_natural": "Естественные ответы (модель сама решает, искать ли)",
        "modes_explicit": "Явно заземлённые в веб ответы",
        "observations": "валидных наблюдений",
        "mentioned": "бренд упомянут",
        "cited": "сайт процитирован как источник",
        "unavailable": "наблюдений недоступно",
        "coverage": "Покрытие провайдеров",
        "competitors": "Самые цитируемые домены в наблюдениях",
        "evidence": "Свидетельства наблюдений (фрагменты)",
        "deliverable_yes": (
            "Прогон достиг порога покрытия: использована одна проверка "
            "AI-присутствия."
        ),
        "deliverable_no": (
            "Прогон не достиг порога покрытия: проверка AI-присутствия не "
            "списана."
        ),
        "searched": "искал в вебе",
        "no_search": "ответил без поиска",
        "status_prefix": "статус",
        "marks": (
            "Названия и знаки провайдеров принадлежат их владельцам. "
            "Crawldy — независимый сервис, не аффилированный с показанными "
            "провайдерами и не одобренный ими."
        ),
        "no_observations": (
            "Для этого провайдера не завершено ни одного валидного "
            "наблюдения; его результат подаётся как недоступный, а не как "
            "нулевое присутствие."
        ),
    },
}


def strings_for(lang: str) -> dict[str, str]:
    """Каталог присутності обраною мовою; невідома — англійська.

    Публічний доступ, бо загальний документ присутності бере звідси ті
    рядки, які мусять бути ОДНАКОВІ в обох видах: застереження про
    торгові назви й опис методу. Два переклади одного юридичного
    речення розійшлися б — питання лише коли.
    """
    return _STRINGS.get(lang, _STRINGS["en"])


def filename_for(url: str, lang: str, *, kind: str = "technical") -> str:
    """Назва файлу: сайт, вид, мова — як у документах аудиту.

    Ідентифікатора прогону тут немає навмисно: у теці завантажень він
    нічого не каже читачеві, а домен каже все. Двом прогонам одного
    сайту одного дня доведеться розрізнятися датою файлу — те саме
    правило, що й для аудиту.
    """
    host = url.replace("https://", "").replace("http://", "").strip("/").split("/")[0]
    return f"ai-presence-{host or 'site'}-{kind}-{lang}.docx"


def render_presence_docx(
    *,
    run: dict[str, Any],
    metrics: dict[str, Any],
    coverage: dict[str, Any],
    observations: list[dict[str, Any]],
    lang: str = "en",
) -> bytes:
    t = _STRINGS.get(lang, _STRINGS["en"])
    document = Document()
    # Та сама типографіка, що в решті клієнтських документів: цей теж
    # їде клієнту, і виглядати він має як частина одного продукту.
    style_document(document)
    host = str(run.get("url") or "").replace("https://", "").replace("http://", "").strip("/")
    add_footer_for(document, host)

    document.add_heading(t["title"], level=0)
    kind = document.add_paragraph()
    kind_run = kind.add_run(label("kind_technical", parse(lang) or DEFAULT_LOCALE))
    kind_run.font.size = Pt(11)
    kind_run.font.color.rgb = MUTED
    document.add_paragraph(f"{t['for']}: {run.get('url', '')}")
    document.add_paragraph(f"{t['brand']}: {run.get('brand', '')}")
    finished = run.get("finished_at") or run.get("created_at")
    if isinstance(finished, datetime):
        document.add_paragraph(
            f"{t['generated']}: {finished.strftime('%Y-%m-%d %H:%M UTC')}"
        )
    method = document.add_paragraph(t["method"])
    method.runs[0].font.size = Pt(9)

    settled = run.get("deliverable")
    if settled is not None:
        document.add_paragraph(
            t["deliverable_yes"] if settled else t["deliverable_no"]
        )

    # Покриття провайдерів у режимах: «x із N», ніколи не голий відсоток.
    document.add_heading(t["coverage"], level=1)
    modes = metrics.get("modes") or {}
    for mode_key, label_key in (("natural", "modes_natural"),
                                ("explicit", "modes_explicit")):
        block = modes.get(mode_key)
        if not block:
            continue
        document.add_heading(t[label_key], level=2)
        for provider, values in sorted((block.get("per_provider") or {}).items()):
            mention = values.get("mention_presence") or {}
            citation = values.get("citation_presence") or {}
            denom = mention.get("denominator") or 0
            unavailable = mention.get("unavailable") or 0
            line = document.add_paragraph(style="List Bullet")
            line.add_run(PROVIDER_NAMES.get(provider, provider)).bold = True
            if denom == 0:
                # Правда про часткове покриття: недоступний провайдер —
                # не нульова присутність.
                line.add_run(f" — {t['no_observations']}")
                continue
            line.add_run(
                f" — {denom} {t['observations']}; "
                f"{t['mentioned']}: {mention.get('numerator', 0)}/{denom}; "
                f"{t['cited']}: {citation.get('numerator', 0)}/"
                f"{citation.get('denominator', 0)}"
                + (f"; {t['unavailable']}: {unavailable}" if unavailable else "")
            )

    observed = metrics.get("observed_domains") or []
    if observed:
        document.add_heading(t["competitors"], level=1)
        for row in observed[:10]:
            document.add_paragraph(
                f"{row.get('domain')} — {row.get('citations')}",
                style="List Bullet",
            )

    document.add_heading(t["evidence"], level=1)
    shown = 0
    for o in observations:
        if o.get("status") != "ok" or shown >= 12:
            continue
        shown += 1
        head = document.add_paragraph()
        head.add_run(
            f"{PROVIDER_NAMES.get(str(o.get('provider')), o.get('provider'))}"
            f" · {o.get('mode', 'natural')} · "
            + (t["searched"] if o.get("searched") else t["no_search"])
        ).bold = True
        document.add_paragraph(str(o.get("prompt_text") or ""))
        excerpt = document.add_paragraph(str(o.get("answer_excerpt") or "")[:400])
        excerpt.runs[0].font.size = Pt(9)

    marks = document.add_paragraph(t["marks"])
    marks.runs[0].font.size = Pt(8)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
